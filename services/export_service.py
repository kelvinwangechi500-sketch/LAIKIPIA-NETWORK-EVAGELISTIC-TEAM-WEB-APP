"""Export Service - PDF, DOCX, Excel generation"""
import io, json
from datetime import datetime

def export_minutes_pdf(minutes) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=18, spaceAfter=6)
        h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=13, spaceAfter=4, textColor=colors.HexColor('#7F1D1D'))
        body = styles['Normal']
        body.fontSize = 11
        body.leading = 16

        story.append(Paragraph("Network Evangelistic Team", styles['Title']))
        story.append(Paragraph(f"Meeting Minutes — {minutes.title}", title_style))
        story.append(Paragraph(f"Date: {minutes.meeting_date.strftime('%d %B %Y')}", body))
        if minutes.recorder:
            story.append(Paragraph(f"Recorded by: {minutes.recorder.full_name}", body))
        story.append(Spacer(1, 0.4*cm))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#B91C1C')))
        story.append(Spacer(1, 0.4*cm))

        if minutes.summary:
            story.append(Paragraph("Summary", h2))
            story.append(Paragraph(minutes.summary.replace('\n','<br/>'), body))
            story.append(Spacer(1, 0.3*cm))

        if minutes.action_points:
            story.append(Paragraph("Action Points", h2))
            try:
                points = json.loads(minutes.action_points)
                for i, pt in enumerate(points, 1):
                    story.append(Paragraph(f"{i}. {pt}", body))
            except:
                story.append(Paragraph(minutes.action_points.replace('\n','<br/>'), body))
            story.append(Spacer(1, 0.3*cm))

        if minutes.transcript:
            story.append(Paragraph("Full Transcript", h2))
            story.append(Paragraph(minutes.transcript.replace('\n','<br/>'), body))

        doc.build(story)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        return f"PDF generation error: {e}".encode()

def export_minutes_docx(minutes) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import json

        doc = Document()
        doc.add_heading('Network Evangelistic Team', 0)
        doc.add_heading(f'Meeting Minutes — {minutes.title}', 1)
        p = doc.add_paragraph()
        p.add_run(f'Date: ').bold = True
        p.add_run(minutes.meeting_date.strftime('%d %B %Y'))
        if minutes.recorder:
            p2 = doc.add_paragraph()
            p2.add_run('Recorded by: ').bold = True
            p2.add_run(minutes.recorder.full_name)

        if minutes.summary:
            doc.add_heading('Summary', 2)
            doc.add_paragraph(minutes.summary)

        if minutes.action_points:
            doc.add_heading('Action Points', 2)
            try:
                points = json.loads(minutes.action_points)
                for pt in points:
                    doc.add_paragraph(pt, style='List Number')
            except:
                doc.add_paragraph(minutes.action_points)

        if minutes.transcript:
            doc.add_heading('Full Transcript', 2)
            doc.add_paragraph(minutes.transcript)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        return f"DOCX generation error: {e}".encode()

def export_attendance_excel(records) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance"

        headers = ["#","Meeting Date","Member Name","Residence","Year","Status","Source","Recorded By","Timestamp"]
        red_fill = PatternFill("solid", fgColor="B91C1C")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = red_fill
            cell.alignment = Alignment(horizontal="center")

        for i, r in enumerate(records, 2):
            recorder_name = r.recorder.full_name if r.recorder else "System"
            ws.append([
                i-1,
                str(r.meeting_date),
                r.user.full_name,
                getattr(r.user,'residence','') or '',
                getattr(r.user,'year','') or '',
                r.status,
                r.source,
                recorder_name,
                r.created_at.strftime("%Y-%m-%d %H:%M"),
            ])
            status_cell = ws.cell(row=i, column=6)
            if r.status == 'present':
                status_cell.font = Font(color="065F46")
            elif r.status == 'absent':
                status_cell.font = Font(color="7F1D1D")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        return f"Excel error: {e}".encode()

def _attendance_pdf(records) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        import io

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        story = [Paragraph("Attendance Report — NET Platform", styles['Title']), Spacer(1, 12)]

        data = [["#","Date","Member","Residence","Year","Status","Source","Recorded By"]]
        for i, r in enumerate(records, 1):
            data.append([
                str(i), str(r.meeting_date), r.user.full_name,
                getattr(r.user,'residence','') or '',
                getattr(r.user,'year','') or '',
                r.status, r.source,
                r.recorder.full_name if r.recorder else 'System'
            ])

        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0), colors.HexColor('#B91C1C')),
            ('TEXTCOLOR',(0,0),(-1,0), colors.white),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
            ('FONTSIZE',(0,0),(-1,-1),8),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, colors.HexColor('#FEF2F2')]),
            ('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#E5E7EB')),
            ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ]))
        story.append(t)
        doc.build(story)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        return f"PDF error: {e}".encode()
